"""Regenerate the branded Word deliverables in senedler/word/ from the markdown
sources in senedler/qaydalar, senedler/roadmap and senedler/en.

    pip install python-docx
    python senedler/tools/make-docx.py

Edit the markdown, run this, commit both. Bump the version line below when the
content changes in a way clients should notice.
"""
import io
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

sys.stdout.reconfigure(encoding="utf-8")

ROOT = "c:/Users/Asus/Desktop/AB\u015E servis/only front"
BRAND = "Kestridge AI"
LEGAL = "Kestridge AI Inc., Illinois, USA"

TEAL = RGBColor(0x0B, 0x7A, 0x67)
INK = RGBColor(0x14, 0x18, 0x1F)
BODY = RGBColor(0x22, 0x26, 0x2E)
FAINT = RGBColor(0x8A, 0x91, 0x9C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SANS = "Segoe UI"
MONO = "Consolas"


def shade(par, fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), fill)
    par._p.get_or_add_pPr().append(el)


def bottom_border(par, color, sz, space):
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    pbdr.append(b)
    par._p.get_or_add_pPr().append(pbdr)


def keep_next(par):
    el = OxmlElement("w:keepNext")
    par._p.get_or_add_pPr().insert(0, el)


def run(par, text, *, font=SANS, size=10.5, bold=False, color=BODY):
    r = par.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold or None
    r.font.color.rgb = color
    return r


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")


def add_inline(par, text, *, size=10.5, base=BODY, strong=INK):
    """Render **bold**, `code` and [links](url) into one paragraph."""
    text = LINK.sub(r"\1", text)
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run(par, piece[2:-2], size=size, bold=True, color=strong)
        elif piece.startswith("`") and piece.endswith("`"):
            run(par, piece[1:-1], font=MONO, size=size, color=strong)
        else:
            run(par, piece, size=size, color=base)


def spacing(par, before=None, after=None):
    if before is not None:
        par.paragraph_format.space_before = Pt(before)
    if after is not None:
        par.paragraph_format.space_after = Pt(after)


# ---------------------------------------------------------------- markdown

def read_blocks(path):
    """Turn a markdown file into a flat list of (kind, payload) blocks."""
    lines = io.open(path, encoding="utf-8").read().split("\n")
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        s = raw.strip()

        if not s:
            i += 1
            continue

        if s.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", code))
            continue

        if re.match(r"^-{3,}$", s):
            i += 1
            continue

        if s.startswith("#### "):
            blocks.append(("h4", s[5:].strip()))
            i += 1
            continue
        if s.startswith("### "):
            blocks.append(("h3", s[4:].strip()))
            i += 1
            continue
        if s.startswith("## "):
            blocks.append(("h2", s[3:].strip()))
            i += 1
            continue
        if s.startswith("# "):
            blocks.append(("h1", s[2:].strip()))
            i += 1
            continue

        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue

        if s.startswith("> "):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(x for x in buf if x)))
            continue

        m = re.match(r"^[-*] \[( |x|X)\] (.*)$", s)
        if m:
            buf = [m.group(2)]
            i += 1
            while i < n and lines[i].startswith("  ") and lines[i].strip() and not re.match(r"^\s*[-*] ", lines[i]):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(("check", " ".join(buf)))
            continue

        m = re.match(r"^[-*] (.*)$", s)
        if m:
            buf = [m.group(1)]
            i += 1
            while i < n and lines[i].startswith("  ") and lines[i].strip() and not re.match(r"^\s*[-*] ", lines[i]) and not lines[i].strip().startswith("|"):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(("bullet", " ".join(buf)))
            continue

        m = re.match(r"^(\d+)\. (.*)$", s)
        if m:
            buf = [m.group(2)]
            i += 1
            while i < n and lines[i].startswith("   ") and lines[i].strip():
                buf.append(lines[i].strip())
                i += 1
            blocks.append(("olist", (m.group(1), " ".join(buf))))
            continue

        buf = [s]
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#|\||>|```|-{3,}|[-*] |\d+\. )", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(("para", " ".join(buf)))
    return blocks


# ---------------------------------------------------------------- rendering

def build(path, out, kicker, meta):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.866)
    sec.top_margin = sec.bottom_margin = Inches(0.7875)

    normal = doc.styles["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(10.5)

    p = doc.add_paragraph()
    spacing(p, after=1)
    run(p, " ".join(BRAND.upper()), size=10, bold=True, color=TEAL)

    p = doc.add_paragraph()
    spacing(p, after=1)
    run(p, kicker, size=8, color=FAINT)

    p = doc.add_paragraph()
    spacing(p, after=6)
    bottom_border(p, "0B7A67", 16, 6)
    run(p, meta, size=8.5, color=FAINT)

    blocks = read_blocks(path)
    for kind, payload in blocks:
        if kind == "h1":
            p = doc.add_paragraph()
            spacing(p, before=10, after=6)
            run(p, payload, size=19, bold=True, color=INK)

        elif kind == "h2":
            p = doc.add_paragraph()
            spacing(p, before=14, after=5)
            keep_next(p)
            bottom_border(p, "D7DBE0", 6, 3)
            run(p, payload, size=13, bold=True, color=TEAL)

        elif kind in ("h3", "h4"):
            p = doc.add_paragraph()
            spacing(p, before=10, after=3)
            keep_next(p)
            run(p, payload, size=11, bold=True, color=INK)

        elif kind == "para":
            p = doc.add_paragraph()
            spacing(p, after=4)
            add_inline(p, payload)

        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            spacing(p, after=4)
            add_inline(p, payload)

        elif kind == "check":
            p = doc.add_paragraph()
            spacing(p, after=2)
            p.paragraph_format.left_indent = Pt(11.35)
            run(p, "\u2610  ", size=11, color=FAINT)
            add_inline(p, payload, size=11, base=FAINT, strong=INK)

        elif kind == "olist":
            num, text = payload
            p = doc.add_paragraph()
            spacing(p, after=4)
            p.paragraph_format.left_indent = Pt(17)
            run(p, "%s.  " % num, size=10.5, bold=True, color=TEAL)
            add_inline(p, text)

        elif kind == "quote":
            p = doc.add_paragraph()
            spacing(p, before=6, after=8)
            p.paragraph_format.left_indent = Pt(11.35)
            shade(p, "EEF4F2")
            add_inline(p, payload, size=10)

        elif kind == "code":
            for line in payload:
                p = doc.add_paragraph()
                spacing(p, after=0)
                p.paragraph_format.left_indent = Pt(8.5)
                shade(p, "F4F6F8")
                run(p, line, font=MONO, size=9, color=INK)
            doc.add_paragraph()

        elif kind == "table":
            rows = payload
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=cols)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tblPr = t._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "6")
                e.set(qn("w:space"), "0")
                e.set(qn("w:color"), "D7DBE0")
                borders.append(e)
            tblPr.append(borders)
            for ri, cells in enumerate(rows):
                row = t.add_row()
                for ci in range(cols):
                    cell = row.cells[ci]
                    cell.paragraphs[0].text = ""
                    par = cell.paragraphs[0]
                    text = cells[ci] if ci < len(cells) else ""
                    if ri == 0:
                        shade(par, "0B7A67")
                        run(par, LINK.sub(r"\1", text).replace("**", ""),
                            size=9.5, bold=True, color=WHITE)
                    else:
                        spacing(par, after=2)
                        add_inline(par, text, size=9.5)
            doc.add_paragraph()

    doc.save(out)
    return out


AREA_AZ = {
    "00-umumi": "Umumi",
    "ai-solutions": "AI Solutions",
    "analytics": "Analytics",
    "automation": "Automation",
    "it-security": "IT Security",
}
AREA_EN = dict(AREA_AZ, **{"00-umumi": "General"})

JOBS = []
for slug, area in AREA_AZ.items():
    JOBS.append(("senedler/qaydalar/%s.md" % slug,
                 "senedler/word/az/%s Is Qaydalari - %s (AZ).docx" % (BRAND, area),
                 "DAXİLİ SƏNƏD · İŞ QAYDALARI",
                 "Versiya 1.1 · Avqust 2026 · " + LEGAL))
    if slug != "00-umumi":
        JOBS.append(("senedler/roadmap/%s.md" % slug,
                     "senedler/word/az/%s Roadmap - %s (AZ).docx" % (BRAND, area),
                     "DAXİLİ SƏNƏD · LAYİHƏ YOL XƏRİTƏSİ",
                     "Versiya 1.1 · Avqust 2026 · " + LEGAL))
for slug, area in AREA_EN.items():
    JOBS.append(("senedler/en/qaydalar/%s.md" % slug,
                 "senedler/word/en/%s Delivery Rules - %s (EN).docx" % (BRAND, area),
                 "INTERNAL DOCUMENT · DELIVERY RULES",
                 "Version 1.1 · August 2026 · " + LEGAL))
    if slug != "00-umumi":
        JOBS.append(("senedler/en/roadmap/%s.md" % slug,
                     "senedler/word/en/%s Roadmap - %s (EN).docx" % (BRAND, area),
                     "INTERNAL DOCUMENT · PROJECT ROADMAP",
                     "Version 1.1 · August 2026 · " + LEGAL))

os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
made = 0
for src, dst, kicker, meta in JOBS:
    if not os.path.exists(src):
        print("MISSING SOURCE", src)
        continue
    build(src, dst, kicker, meta)
    made += 1
    print("wrote", dst)
print("total", made)
